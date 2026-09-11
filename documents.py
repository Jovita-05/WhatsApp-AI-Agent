import re
from pathlib import Path

from docx import Document

import io
import os
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

import sharepoint




# Create Word document locally
def create_document(
    request_text,
    content,
):
    words = re.findall(
        r"[A-Za-z0-9]+",
        request_text,
    )

    short_name = "_".join(
        words[:6]
    ) or "generated_document"

    filename = (
        f"{short_name[:60]}.docx"
    )

    document = Document()
    document.add_heading(
        "Generated Document",
        level=1,
    )

    for paragraph in str(content).splitlines():
        if paragraph.strip():
            document.add_paragraph(
                paragraph.strip()
            )

    document.save(filename)

    return filename


# Read approved SharePoint documents
def read_sharepoint():
    headers, folders, _ = (
        sharepoint.connect_sharepoint()
    )

    text = []

    for folder_name in (
        "Policies",
        "Reports",
    ):
        folder = sharepoint.find_folder(
            folders,
            folder_name,
        )

        if not folder:
            continue

        drive_id, files = (
            sharepoint.get_folder_files(
                headers,
                folder,
            )
        )

        for item in files:
            name = item.get(
                "name",
                "",
            )

            lower_name = (
                name.lower()
            )

            if not lower_name.endswith((
                ".docx",
                ".xlsx",
            )):
                continue

            content = (
                sharepoint.download_drive_item(
                    headers,
                    drive_id,
                    item["id"],
                )
            )

            text.append(
                f"\n--- {name} ---"
            )

            if lower_name.endswith(
                ".docx"
            ):
                document = Document(
                    io.BytesIO(content)
                )

                text.extend(
                    paragraph.text
                    for paragraph
                    in document.paragraphs
                    if paragraph.text.strip()
                )

            else:
                workbook = load_workbook(
                    io.BytesIO(content),
                    data_only=True,
                    read_only=True,
                )

                try:
                    for sheet in workbook:
                        for row in (
                            sheet.iter_rows(
                                values_only=True
                            )
                        ):
                            text.append(
                                str(row)
                            )

                finally:
                    workbook.close()

    return "\n".join(text)


# Upload bytes to SharePoint folder
def upload_bytes_to_folder(
    headers,
    folder,
    filename,
    content,
    mime_type,
):
    filename = os.path.basename(
        filename
    )

    drive_id = (
        folder[
            "parentReference"
        ]["driveId"]
    )

    folder_id = folder["id"]

    response = sharepoint.session.put(
        f"{sharepoint.GRAPH_URL}/drives/"
        f"{drive_id}/items/"
        f"{folder_id}:/{filename}:/content",
        headers={
            **headers,
            "Content-Type":
                mime_type,
        },
        data=content,
        timeout=120,
    )

    response.raise_for_status()

    return response.json()


# Upload generated document
def upload_document(
    filename
):
    headers, folders, _ = (
        sharepoint.connect_sharepoint()
    )

    folder = sharepoint.find_folder(
        folders,
        "Generated Documents",
    )

    if not folder:
        raise RuntimeError(
            "Generated Documents folder "
            "was not found."
        )

    with open(
        filename,
        "rb",
    ) as file:
        return upload_bytes_to_folder(
            headers,
            folder,
            os.path.basename(
                filename
            ),
            file.read(),
            "application/octet-stream",
        )


# Upload WhatsApp document
def upload_whatsapp_document(
    filename,
    file_content,
    mime_type=
        "application/octet-stream",
):
    headers, folders, _ = (
        sharepoint.connect_sharepoint()
    )

    folder = sharepoint.find_folder(
        folders,
        "WhatsApp Uploads",
    )

    if not folder:
        raise RuntimeError(
            "WhatsApp Uploads folder "
            "was not found."
        )

    return upload_bytes_to_folder(
        headers,
        folder,
        filename,
        file_content,
        mime_type,
    )


# Extract readable text from Word or Excel documents
def extract_document_text(
    filename,
    file_content,
):
    filename_lower = (
        filename.lower()
    )
#docx
    if filename_lower.endswith(
        ".docx"
    ):
        document = Document(
            io.BytesIO(
                file_content
            )
        )

        text = [
            paragraph.text
            for paragraph
            in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in (
            document.tables
        ):
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell
                    in row.cells
                    if cell.text.strip()
                ]

                if values:
                    text.append(
                        " | ".join(
                            values
                        )
                    )

        if text:
            return "\n".join(
                text
            )

        # Fallback for Word XML
        with zipfile.ZipFile(
            io.BytesIO(
                file_content
            )
        ) as archive:
            xml_data = (
                archive.read(
                    "word/document.xml"
                )
            )

        root = ET.fromstring(
            xml_data
        )

        namespace = {
            "w":
                "http://schemas.openxmlformats.org/"
                "wordprocessingml/2006/main"
        }

        return "\n".join(
            node.text
            for node
            in root.findall(
                ".//w:t",
                namespace,
            )
            if node.text
            and node.text.strip()
        )
#excel
    if filename_lower.endswith(
        ".xlsx"
    ):
        workbook = load_workbook(
            io.BytesIO(
                file_content
            ),
            data_only=True,
            read_only=True,
        )

        text = []

        try:
            for sheet in workbook:
                for row in (
                    sheet.iter_rows(
                        values_only=True
                    )
                ):
                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:
                        text.append(
                            " | ".join(
                                values
                            )
                        )

        finally:
            workbook.close()

        return "\n".join(
            text
        )

    raise ValueError(
        "Only .docx and .xlsx files "
        "can be summarized."
    )
