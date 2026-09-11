import io
import logging
import os
import time
import requests

from dotenv import load_dotenv
from openpyxl import load_workbook


load_dotenv()


TENANT_ID = os.getenv("SHAREPOINT_TENANT_ID", "").strip()
CLIENT_ID = os.getenv("SHAREPOINT_CLIENT_ID", "").strip()
CLIENT_SECRET = os.getenv("SHAREPOINT_CLIENT_SECRET", "").strip()
HOSTNAME = os.getenv(
    "SHAREPOINT_HOSTNAME", "softflowcloud.sharepoint.com"
).strip()
SITE_PATH = os.getenv(
    "SHAREPOINT_SITE_PATH", "/sites/TrainingCenter/JovitaTest"
).strip()

GRAPH_URL = "https://graph.microsoft.com/v1.0"

log = logging.getLogger(__name__)

session = requests.Session()

# raise_for_status only ever tells us the status code. Microsoft puts the useful part in the body, so log it here once and every Graph call in the app gets it for free.
def log_failed_response(response, *args, **kwargs):
    if not response.ok:
        log.error(
            "Graph %s %s -> %s %s",
            response.request.method,
            response.url,
            response.status_code,
            response.text[:1000],
        )

    return response


session.hooks["response"].append(log_failed_response)

# The token lasts an hour and the site ID never changes, so hold on to both rather than asking Microsoft again on every request.
_cached_token = None
_cached_token_expires_at = 0
_cached_site_id = None
_cached_list_ids = {}

# Get a Graph API access token, reusing it until shortly before it expires
def get_access_token():
    global _cached_token, _cached_token_expires_at

    if _cached_token and time.time() < _cached_token_expires_at:
        return _cached_token

    if not all([TENANT_ID, CLIENT_ID, CLIENT_SECRET]):
        raise ValueError("SharePoint credentials are missing.")

    response = session.post(
        f"https://login.microsoftonline.com/{TENANT_ID}/oauth2/v2.0/token",
        data={
            "client_id": CLIENT_ID,
            "client_secret": CLIENT_SECRET,
            "scope": "https://graph.microsoft.com/.default",
            "grant_type": "client_credentials",
        },
        timeout=30,
    )
    response.raise_for_status()
    token_data = response.json()

    _cached_token = token_data["access_token"]
    # Expire a minute early so a request that is already in flight cannot
    # arrive with a dead token.
    _cached_token_expires_at = (
        time.time() + token_data.get("expires_in", 3600) - 60
    )

    return _cached_token

def get_graph_context():
    global _cached_site_id

    headers = {"Authorization": f"Bearer {get_access_token()}"}

    if _cached_site_id is None:
        response = session.get(
            f"{GRAPH_URL}/sites/{HOSTNAME}:{SITE_PATH}",
            headers=headers,
            timeout=30,
        )
        response.raise_for_status()
        _cached_site_id = response.json()["id"]

    return headers, _cached_site_id

# Connect to SharePoint and retrieve root folders
def connect_sharepoint():
    headers, site_id = get_graph_context()

    response = session.get(
        f"{GRAPH_URL}/sites/{site_id}/drive/root/children",
        headers=headers,
        timeout=30,
    )
    response.raise_for_status()

    return headers, response.json().get("value", []), site_id

# Folder names in SharePoint are not reliably capitalised, hence casefold.
def find_folder(folders, folder_name):
    target = folder_name.casefold()

    return next(
        (f for f in folders if f.get("name", "").casefold() == target),
        None,
    )

# Get files from SharePoint folder
def get_folder_files(headers, folder):
    drive_id = folder["parentReference"]["driveId"]

    response = session.get(
        f"{GRAPH_URL}/drives/{drive_id}/items/{folder['id']}/children",
        headers=headers,
        timeout=30,
    )
    response.raise_for_status()

    return drive_id, response.json().get("value", [])

# Download SharePoint file
def download_drive_item(headers, drive_id, item_id):
    response = session.get(
        f"{GRAPH_URL}/drives/{drive_id}/items/{item_id}/content",
        headers=headers,
        timeout=60,
    )
    response.raise_for_status()

    return response.content

# Find SharePoint list ID
def get_list_id(headers, site_id, list_name):
    response = session.get(
        f"{GRAPH_URL}/sites/{site_id}/lists",
        headers=headers,
        timeout=30,
    )
    response.raise_for_status()

    target = list_name.casefold()

    for sp_list in response.json().get("value", []):
        if sp_list.get("displayName", "").casefold() == target:
            return sp_list["id"]

    return None

# Look up a list by its display name once, then remember its ID.
def get_list_context(list_name):
    headers, site_id = get_graph_context()

    if list_name not in _cached_list_ids:
        list_id = get_list_id(headers, site_id, list_name)

        if not list_id:
            raise RuntimeError(f"{list_name} list was not found.")

        _cached_list_ids[list_name] = list_id

    return headers, site_id, _cached_list_ids[list_name]

def list_items_url(site_id, list_id, item_id=""):
    url = f"{GRAPH_URL}/sites/{site_id}/lists/{list_id}/items"

    return f"{url}/{item_id}" if item_id else url

# JSON body plus the auth headers, for POST and PATCH.
def json_headers(headers):
    return {**headers, "Content-Type": "application/json"}

def normalize_number(number):
    return str(number).replace("+", "").strip()

# Salary lookups are keyed on the sender's WhatsApp number, so nobody can
def get_my_salary(sender_number):
    sender_number = normalize_number(sender_number)
    not_found = (
        "I could not find a salary record for your WhatsApp number."
    )

    headers, folders, _ = connect_sharepoint()

    salary_folder = next(
        (f for f in folders if "salar" in f.get("name", "").lower()),
        None,
    )

    if not salary_folder:
        return not_found

    drive_id, files = get_folder_files(headers, salary_folder)

    for item in files:
        if not item.get("name", "").lower().endswith(".xlsx"):
            continue

        content = download_drive_item(headers, drive_id, item["id"])
        workbook = load_workbook(
            io.BytesIO(content), data_only=True, read_only=True
        )

        try:
            rows = workbook.active.iter_rows(values_only=True)
            columns = next(rows, None)

            if not columns:
                continue

            for row in rows:
                employee = dict(zip(columns, row))

                if normalize_number(
                    employee.get("WhatsApp Number", "")
                ) == sender_number:
                    return (
                        f"Your salary is {employee.get('Salary')} "
                        f"{employee.get('Currency')}."
                    )
        finally:
            workbook.close()

    return not_found